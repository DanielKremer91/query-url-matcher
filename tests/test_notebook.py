import ast
import json
import re
from pathlib import Path

import pytest

from qum import colab
from qum.embeddings.base import EmbeddingError
from qum.ingest import IngestError
from qum.models import MODELS
from tools.build_notebook import CELLS, build

ROOT = Path(__file__).resolve().parents[1]


def _code_cells():
    return [source for kind, source in CELLS if kind == "code"]


def test_committed_notebook_matches_builder():
    committed = json.loads((ROOT / "query_url_matcher.ipynb").read_text(encoding="utf-8"))
    assert committed == build()


def test_code_cells_are_valid_python():
    for source in _code_cells():
        python = "\n".join(line for line in source.splitlines() if not line.lstrip().startswith(("!", "%")))
        ast.parse(python)


def test_every_code_cell_is_a_titled_form():
    for source in _code_cells():
        assert source.startswith("#@title "), source[:40]


def test_model_dropdown_lists_exactly_the_registry_labels():
    source = next(s for s in _code_cells() if "modell = " in s)
    options = json.loads(re.search(r"modell = .*#@param (\[.*\])", source).group(1))
    assert options == [spec.label for spec in MODELS.values()]


def test_notebook_has_its_steps_in_order():
    titles = [re.match(r"#@title (.*?)( \{|$)", s.splitlines()[0]).group(1) for s in _code_cells()]
    expected = ["2", "3", "4", "5", "6", "7a", "7b", "7c", "8"]
    assert [t.split(":")[0].split(" (")[0] for t in titles] == [f"Schritt {n}" for n in expected]
    assert any(kind == "markdown" and "Schritt 1" in source for kind, source in CELLS)


def test_no_keys_and_no_retired_model_in_notebook():
    text = (ROOT / "query_url_matcher.ipynb").read_text(encoding="utf-8")
    assert "sk-" not in text
    assert "text-embedding-004" not in text


def test_guard_turns_known_errors_into_notebook_stop(capsys):
    for error in (IngestError("Spalte fehlt"), EmbeddingError("Key fehlt")):
        with pytest.raises(colab.NotebookStop):
            with colab.guard():
                raise error
    assert "❌ Spalte fehlt" in capsys.readouterr().out


def test_guard_lets_other_errors_through():
    with pytest.raises(ZeroDivisionError):
        with colab.guard():
            1 / 0


def test_notebook_stop_hides_traceback():
    assert colab.NotebookStop()._render_traceback_() == []


def test_require_stops_with_german_hint_for_missing_or_none(capsys):
    colab.require({"result": 1}, 4, "result")
    for namespace in ({}, {"result": None}):
        with pytest.raises(colab.NotebookStop):
            colab.require(namespace, 4, "result")
    assert "❌ Bitte zuerst Schritt 4 (Matching) ausführen." in capsys.readouterr().out


def test_require_accepts_sub_steps(capsys):
    for step, text in (("7a", "Vorschläge ansehen"), ("7b", "Schwelle festlegen"), ("7c", "Urteile bilden")):
        with pytest.raises(colab.NotebookStop):
            colab.require({}, step, "x")
        assert f"❌ Bitte zuerst Schritt {step} ({text}) ausführen." in capsys.readouterr().out


def test_invalidate_sets_names_to_none():
    namespace = {"decisions": 1, "top": 2}
    colab.invalidate(namespace, "decisions", "settings")
    assert namespace == {"decisions": None, "top": 2, "settings": None}


def _fake_userdata(monkeypatch, get):
    import sys
    import types

    class SecretNotFoundError(Exception):
        pass

    class NotebookAccessError(Exception):
        pass

    userdata = types.ModuleType("google.colab.userdata")
    userdata.SecretNotFoundError = SecretNotFoundError
    userdata.NotebookAccessError = NotebookAccessError
    userdata.get = get(SecretNotFoundError, NotebookAccessError)
    package = types.ModuleType("google.colab")
    package.userdata = userdata
    monkeypatch.setitem(sys.modules, "google", types.ModuleType("google"))
    monkeypatch.setitem(sys.modules, "google.colab", package)
    monkeypatch.setitem(sys.modules, "google.colab.userdata", userdata)


def test_secret_returns_value_or_none_when_missing(monkeypatch):
    def get(not_found, no_access):
        def inner(name):
            if name == "SET":
                return "key"
            raise not_found(name)

        return inner

    _fake_userdata(monkeypatch, get)
    assert colab.secret("SET") == "key"
    assert colab.secret("OTHER") is None


def test_secret_without_notebook_access_explains_the_switch(monkeypatch, capsys):
    def get(not_found, no_access):
        def inner(name):
            raise no_access(name)

        return inner

    _fake_userdata(monkeypatch, get)
    with pytest.raises(colab.NotebookStop):
        colab.secret("OPENAI_API_KEY")
    out = capsys.readouterr().out
    assert "OPENAI_API_KEY" in out and "Notebook-Zugriff" in out


@pytest.mark.parametrize("name, data", [("kaputt.xlsx", b"das ist keine Excel-Datei"), ("kaputt.csv", b'a,b\n"1,2\n3')])
def test_read_table_turns_parse_errors_into_german_hint(name, data):
    with pytest.raises(IngestError, match=f"'{name}' konnte nicht gelesen werden"):
        colab.read_table(data, name)


def test_read_table_passes_ingest_errors_through():
    with pytest.raises(IngestError, match="nicht unterstützt"):
        colab.read_table(b"x", "bild.png")


# --- Rauchtest: die Zellen der Schritte 3 bis 8 laufen in einem gemeinsamen Namensraum ---

EXAMPLES = ROOT / "examples"
CHOICES = ["Aus Rankings kalibriert", "Mittlerer bester Score (nicht kalibriert)", "Eigener Wert"]
CALIBRATED, MEDIAN, OWN = CHOICES
NEED_7A = "❌ Bitte zuerst Schritt 7a (Vorschläge ansehen) ausführen."
NEED_7B = "❌ Bitte zuerst Schritt 7b (Schwelle festlegen) ausführen."
NEED_7C = "❌ Bitte zuerst Schritt 7c (Urteile bilden) ausführen."
KEPT = "Bis dahin gelten die Urteile aus dem letzten Lauf von Schritt 7c."


def _source(step) -> str:
    """Quelltext der Zelle zu Schritt 2 bis 8 oder "7a", "7b", "7c"."""
    return next(s for s in _code_cells() if re.match(rf"#@title Schritt {step}[: ]", s))


def _cell_code(step) -> str:
    return "\n".join(line for line in _source(step).splitlines() if not line.lstrip().startswith(("!", "%")))


def _form_names(step) -> set:
    return set(re.findall(r"^(\w+) = .*? #@param", _cell_code(step), flags=re.M))


class Notebook:
    def __init__(self, monkeypatch, tmp_path):
        from qum.embeddings.cache import CachedEmbedder
        from tests.conftest import FAKE_SPEC, FakeEmbedder

        self.uploads, self.downloads = [], []
        monkeypatch.chdir(tmp_path)
        monkeypatch.setattr(colab, "upload", lambda title: self.uploads.pop(0))
        monkeypatch.setattr(colab, "download", lambda path: self.downloads.append(str(path)))
        monkeypatch.setattr(colab, "secret", lambda name: None)
        self.ns = {"display": lambda frame: None, "spec": FAKE_SPEC, "embedder": CachedEmbedder(FakeEmbedder())}

    def run(self, step, *files, **form):
        self.uploads += [(name, (EXAMPLES / name).read_bytes()) for name in files]
        code = _cell_code(step)
        for name, value in form.items():  # Formularwerte ersetzen die Vorgaben der #@param-Zeilen
            code, n = re.subn(rf"^{name} = .*?(?= #@param)", f"{name} = {value!r}", code, flags=re.M)
            assert n == 1, name
        exec(compile(code, f"<schritt {step}>", "exec"), self.ns)


@pytest.fixture
def nb(monkeypatch, tmp_path):
    return Notebook(monkeypatch, tmp_path)


def _load(nb):
    nb.run(3, "queries.csv", "frog_export.csv")
    nb.run(4)


def _step_7(nb, steps, choice, form):
    """Führt die Zellen aus steps aus; jedes Formularfeld geht an die Zelle, in der es steht."""
    form = {"schwelle_waehlen": choice, **form}
    for step in steps:
        names = _form_names(step)
        nb.run(step, **{name: value for name, value in form.items() if name in names})
    unused = set(form) - set().union(*(_form_names(step) for step in steps))
    assert not unused, unused


def _threshold(nb, choice=MEDIAN, **form):
    """Schritt 7a und 7b."""
    _step_7(nb, ("7a", "7b"), choice, form)


def _verdicts(nb, choice=MEDIAN, **form):
    """Schritt 7a, 7b und 7c."""
    _step_7(nb, ("7a", "7b", "7c"), choice, form)


def test_smoke_full_flow(nb, tmp_path, capsys):
    _load(nb)
    nb.run(5, "rankings.csv")
    nb.run(6, "serps.csv")
    _verdicts(nb)
    nb.run(8)
    out = capsys.readouterr().out
    assert "✅ Schritt 8 fertig: 5 Blätter exportiert." in out
    assert nb.downloads == ["query_url_matcher.xlsx"]
    assert _sheet_names(tmp_path / "query_url_matcher.xlsx") == SHEETS
    assert nb.ns["new_pages"] is not None
    settings = nb.ns["settings"]
    assert settings["Eigene Rankings"] == "ja, aus Datei"
    for key in ("Kalibrierung bis Position", "SERP-Überschneidung (%)", "Cluster-Dichte (%)"):
        assert key in settings
    assert "Treffer je Query (Top-N)" not in settings


SHEETS = ["Übersicht", "Kannibalisierungsgefahr", "Potentielle Content-Lücken", "Chunk auf anderer Seite", "Lesehilfe"]


def _sheet_names(path):
    from openpyxl import load_workbook

    return load_workbook(path).sheetnames


def test_smoke_minimal_flow_without_rankings_and_serps(nb, capsys):
    _load(nb)
    nb.run("7a")
    assert "Aus Rankings kalibriert: nicht verfügbar. Es sind keine Rankings geladen (Schritt 5 oder 6)." in capsys.readouterr().out
    nb.run("7b", schwelle_waehlen=MEDIAN)
    nb.run("7c")
    nb.run(8)
    out = capsys.readouterr().out
    assert "per Konstruktion darunter" in out
    assert nb.downloads == ["query_url_matcher.xlsx"]
    assert nb.ns["settings"]["Eigene Rankings"] == "nein"
    assert _sheet_names(nb.downloads[0]) == SHEETS
    assert nb.ns["new_pages"] is None


def test_smoke_serps_alone_derive_own_rankings_and_refresh(nb, capsys):
    _load(nb)
    own = "getreidefreies hundefutter;https://www.tierbedarf.example/hundefutter/getreidefreies-hundefutter;5;Organic\n"
    serps_with_own = (EXAMPLES / "serps.csv").read_bytes() + own.encode()
    nb.uploads.append(("serps.csv", serps_with_own))
    nb.run(6)
    assert nb.ns["rankings_source"] == "SERPs" and len(nb.ns["rankings"]) == 1
    nb.uploads.append(("serps.csv", (EXAMPLES / "serps.csv").read_bytes()))
    nb.run(6)  # neue SERPs ohne eigene URL: die alten abgeleiteten Rankings dürfen nicht bleiben
    assert nb.ns["rankings"] is None and nb.ns["rankings_source"] is None
    nb.uploads.append(("serps.csv", serps_with_own))
    nb.run(6)
    assert nb.ns["rankings_source"] == "SERPs"
    nb.run(5, "rankings.csv")
    nb.run(6, "serps.csv")
    assert nb.ns["rankings_source"] == "Datei"


def test_smoke_step_8_before_step_7c_points_to_step_7c(nb, capsys):
    _load(nb)
    with pytest.raises(colab.NotebookStop):
        nb.run(8)
    assert NEED_7C in capsys.readouterr().out
    _threshold(nb)
    with pytest.raises(colab.NotebookStop):
        nb.run(8)
    assert NEED_7C in capsys.readouterr().out


def test_smoke_step_before_model_or_files_points_back(nb, capsys):
    nb.ns["embedder"] = None
    with pytest.raises(colab.NotebookStop):
        nb.run(3, "queries.csv", "frog_export.csv")
    assert "Bitte zuerst Schritt 2" in capsys.readouterr().out
    nb2 = nb
    nb2.ns["embedder"] = object()
    with pytest.raises(colab.NotebookStop):
        nb2.run(4)
    assert "Bitte zuerst Schritt 3" in capsys.readouterr().out


def test_smoke_rerunning_step_4_invalidates_verdicts(nb, capsys):
    _load(nb)
    _verdicts(nb)
    assert nb.ns["decisions"] is not None
    nb.run(4, bewertungsgrundlage="Kombi")
    for name in ("decisions", "settings", "threshold", "threshold_source", "threshold_label"):
        assert nb.ns[name] is None, name
    with pytest.raises(colab.NotebookStop):
        nb.run(8)
    assert NEED_7C in capsys.readouterr().out
    for step, hint in (("7c", NEED_7B), ("7b", NEED_7A)):
        with pytest.raises(colab.NotebookStop):
            nb.run(step)
        assert hint in capsys.readouterr().out


def test_smoke_new_files_reset_rankings_and_serps_with_note(nb, capsys):
    _load(nb)
    nb.run(5, "rankings.csv")
    nb.run(6, "serps.csv")
    _verdicts(nb)
    capsys.readouterr()
    nb.run(3, "queries.csv", "frog_export.csv")
    out = capsys.readouterr().out
    assert "ℹ️ Rankings und SERPs wurden zurückgesetzt. Führe Schritt 5 und 6 bei Bedarf erneut aus." in out
    assert nb.ns["rankings"] is None and nb.ns["serps"] is None and nb.ns["result"] is None
    for step in ("7a", "7b", "7c"):
        with pytest.raises(colab.NotebookStop):
            nb.run(step)
        assert "Bitte zuerst Schritt 4 (Matching) ausführen." in capsys.readouterr().out


def test_smoke_ranking_file_of_another_project_keeps_plain_matching(nb, capsys):
    _load(nb)
    foreign = "Keyword;URL;Position\nganz anderes thema;https://fremd.example/x;3\n"
    nb.uploads.append(("fremd.csv", foreign.encode()))
    nb.run(5)
    out = capsys.readouterr().out
    assert "⚠️ Keine der Ranking-Zeilen passt zu deinen Queries. Prüfe, ob die Datei zu diesem Projekt gehört." in out
    assert nb.ns["rankings"] is None
    _verdicts(nb)
    assert nb.ns["settings"]["Eigene Rankings"] == "nein"


def test_smoke_ranking_urls_outside_the_export_are_flagged(nb, capsys):
    _load(nb)
    known = "Keyword;URL;Position\ngetreidefreies hundefutter;https://fremd.example/x;3\n"
    nb.uploads.append(("fremd.csv", known.encode()))
    nb.run(5)
    assert "Keine der Ranking-URLs steht im Frog-Export" in capsys.readouterr().out
    assert nb.ns["rankings"] is None


def test_smoke_serps_of_another_project_are_dropped(nb, capsys):
    _load(nb)
    nb.uploads.append(("fremd.csv", b"Keyword;URL;Position;Type\nganz anderes thema;https://x.example/a;1;Organic\n"))
    nb.run(6)
    assert "Kein SERP-Keyword passt zu deinen Queries" in capsys.readouterr().out
    assert nb.ns["serps"] is None


def test_smoke_numeric_inputs_are_validated(nb, capsys):
    nb.run(3, "queries.csv", "frog_export.csv")
    for form, hint in (
        ({"chunk_groesse": -1}, "nicht negativ"),
        ({"chunk_groesse": 10, "chunk_overlap": 10}, "Overlap"),
    ):
        with pytest.raises(colab.NotebookStop):
            nb.run(4, **form)
        assert hint in capsys.readouterr().out
    nb.run(4, chunk_groesse=0, chunk_overlap=0)
    with pytest.raises(colab.NotebookStop):
        nb.run("7a", kalibrierung_bis_position=0)
    assert "mindestens 1" in capsys.readouterr().out
    _threshold(nb)
    for form in ({"rankt_gut_bis_position": 0}, {"sichtbar_bis_position": 0}):
        with pytest.raises(colab.NotebookStop):
            nb.run("7c", **form)
        assert "mindestens 1" in capsys.readouterr().out


def test_smoke_proposals_explain_why_calibration_is_missing(nb, capsys):
    _load(nb)
    nb.run(5, "rankings.csv")
    nb.run("7a")
    out = capsys.readouterr().out
    assert "Aus Rankings kalibriert: nicht verfügbar. Nur 4 Ranking-Paare bis Position 5, für eine Kalibrierung sind 20 nötig." in out
    assert "per Konstruktion darunter" in out


def test_smoke_header_only_ranking_file_stops_with_hint(nb, capsys):
    _load(nb)
    nb.uploads.append(("leer.csv", b"Keyword;URL;Position\n"))
    with pytest.raises(colab.NotebookStop):
        nb.run(5)
    assert "❌ Die Datei enthält keine verwertbaren Zeilen" in capsys.readouterr().out
    assert nb.ns.get("rankings") is None


def _readme_notes(path):
    import pandas as pd

    from qum import labels as L

    readme = pd.read_excel(path, sheet_name="Lesehilfe")
    return readme[readme[L.R_AREA] == "Hinweis"][L.R_TEXT].tolist()


def test_smoke_uncalibrated_threshold_is_flagged_in_counts_and_export(nb, tmp_path, capsys):
    from qum import export

    _load(nb)
    _verdicts(nb)
    assert "potentielle Content-Lücken (Schwelle nicht kalibriert)" in capsys.readouterr().out
    nb.run(8)
    assert export.CAVEAT_THRESHOLD["median"] in _readme_notes(tmp_path / "query_url_matcher.xlsx")


def _zip_header(path):
    import zipfile

    with zipfile.ZipFile(path) as archive:
        assert archive.namelist() == [
            "uebersicht.csv", "kannibalisierungsgefahr.csv", "potentielle_content_luecken.csv", "chunk_auf_anderer_seite.csv",
            "lesehilfe.csv",
        ]
        return archive.read("uebersicht.csv").decode("utf-8-sig").splitlines()[0]


def test_smoke_step_8_csv_zip_uses_semicolon_by_default_and_comma_on_request(nb, tmp_path):
    _load(nb)
    _verdicts(nb)
    nb.run(8, zusaetzlich_csv_zip=True)
    assert ";" in _zip_header(tmp_path / "query_url_matcher_csv.zip")
    nb.run(8, zusaetzlich_csv_zip=True, csv_trennzeichen="Komma")
    header = _zip_header(tmp_path / "query_url_matcher_csv.zip")
    assert "," in header and ";" not in header


def test_step_8_separator_dropdown_defaults_to_semicolon():
    source = next(s for s in _code_cells() if s.startswith("#@title Schritt 8"))
    match = re.search(r'^csv_trennzeichen = (".*?") #@param (\[.*\])$', source, flags=re.M)
    assert json.loads(match.group(1)) == "Semikolon (für deutsches Excel)"
    assert json.loads(match.group(2)) == ["Semikolon (für deutsches Excel)", "Komma"]


def test_smoke_threshold_typed_by_hand_is_named_in_export(nb, tmp_path, capsys):
    from qum import export

    _load(nb)
    _verdicts(nb, OWN, eigene_schwelle=0.5)
    out = capsys.readouterr().out
    assert "potentielle Content-Lücken" in out and "(Schwelle nicht kalibriert)" not in out
    nb.run(8)
    notes = _readme_notes(tmp_path / "query_url_matcher.xlsx")
    assert export.CAVEAT_THRESHOLD["manuell"] in notes and export.CAVEAT_THRESHOLD["median"] not in notes


def test_notebook_does_not_claim_rankings_unlock_cannibalisation_or_gaps():
    text = "\n".join(source for _, source in CELLS)
    assert "Schaltet die Urteile zu Kannibalisierung und Content-Lücke frei" not in text
    assert "Mit eigenen Rankings kommen die Urteile zu Kannibalisierung und Content-Lücke dazu" not in text
    step5 = next(s for s in _code_cells() if s.startswith("#@title Schritt 5"))
    assert "ermöglicht eine Kalibrierung der Schwelle" in step5 and "ab 20 gut rankenden Paaren" in step5
    intro = next(source for kind, source in CELLS if kind == "markdown")
    assert "ermöglichen eine Kalibrierung der Schwelle" in intro and "kalibrieren die Schwelle" not in intro


def test_example_files_use_reserved_example_domains():
    from urllib.parse import urlsplit

    import pandas as pd

    for name, column in (("frog_export.csv", "Address"), ("rankings.csv", "URL"), ("serps.csv", "URL")):
        hosts = pd.read_csv(EXAMPLES / name, sep=";")[column].map(lambda u: urlsplit(u).netloc)
        assert hosts.str.endswith(".example").all(), name
    assert "www.tierbedarf.example" in (EXAMPLES / "frog_export.csv").read_text(encoding="utf-8")


class _FakeCuda:
    def __init__(self, available):
        self.available, self.calls = available, 0

    def is_available(self):
        self.calls += 1
        return self.available


def _run_step_2(nb, monkeypatch, label, gpu):
    import sys
    import types

    import qum.embeddings
    from qum.embeddings.cache import CachedEmbedder
    from tests.conftest import FakeEmbedder

    torch = types.ModuleType("torch")
    torch.cuda = _FakeCuda(gpu)
    monkeypatch.setitem(sys.modules, "torch", torch)
    monkeypatch.setattr(qum.embeddings, "make_embedder", lambda spec, **kw: CachedEmbedder(FakeEmbedder(spec)))
    monkeypatch.setattr(colab, "secret", lambda name: "test-key")
    nb.run(2, modell=label)
    return torch.cuda


def test_smoke_step_2_warns_without_gpu_for_local_models(nb, monkeypatch, capsys):
    cuda = _run_step_2(nb, monkeypatch, MODELS["e5-large"].label, gpu=False)
    out = capsys.readouterr().out
    assert cuda.calls == 1
    assert "⚠️ Colab hat keine GPU zugeteilt" in out and "T4 GPU" in out
    assert "✅ Schritt 2 fertig" in out


def test_smoke_step_2_stays_quiet_with_gpu(nb, monkeypatch, capsys):
    _run_step_2(nb, monkeypatch, MODELS["e5-large"].label, gpu=True)
    assert "keine GPU zugeteilt" not in capsys.readouterr().out


def test_smoke_step_2_does_not_check_gpu_for_api_models(nb, monkeypatch, capsys):
    cuda = _run_step_2(nb, monkeypatch, MODELS["openai"].label, gpu=False)
    assert cuda.calls == 0
    assert "keine GPU zugeteilt" not in capsys.readouterr().out


def test_intro_says_the_gpu_is_requested_automatically():
    intro = next(source for kind, source in CELLS if kind == "markdown")
    assert "automatisch eine kostenlose T4-GPU" in intro


def test_smoke_rejected_rerun_of_step_4_keeps_previous_settings(nb, capsys):
    nb.run(3, "queries.csv", "frog_export.csv")
    nb.run(4)
    before = {name: nb.ns[name] for name in ("size", "overlap", "basis", "result")}
    with pytest.raises(colab.NotebookStop):
        nb.run(4, chunk_groesse=10, chunk_overlap=10, bewertungsgrundlage="Kombi")
    assert {name: nb.ns[name] for name in before} == before


def test_smoke_default_overlap_not_below_chunk_size_explains_itself(nb, capsys):
    nb.run(3, "queries.csv", "frog_export.csv")
    capsys.readouterr()
    with pytest.raises(colab.NotebookStop):
        nb.run(4, chunk_groesse=1, chunk_overlap=0)  # Standard-Overlap des Test-Modells ist 1
    out = capsys.readouterr().out
    assert "Standard-Overlap des Modells (1)" in out and "Trage 0 ein" not in out


def test_smoke_serps_with_few_urls_per_keyword_warn(nb, capsys):
    _load(nb)
    own_only = (
        "Keyword;URL;Position;Type\n"
        "getreidefreies hundefutter;https://www.tierbedarf.example/hundefutter/getreidefreies-hundefutter;5;Organic\n"
    )
    nb.uploads.append(("serps.csv", own_only.encode()))
    nb.run(6)
    assert "⚠️ Im Schnitt nur 1 URLs je Keyword." in capsys.readouterr().out
    nb.run(6, "serps.csv")
    assert "Im Schnitt nur" not in capsys.readouterr().out


def test_step_7c_fine_settings_keep_their_defaults():
    source = _source("7c")
    for line in ("rankt_gut_bis_position = 10", "abstand_fast_gleich = 0.01", "sichtbar_bis_position = 20",
                 "luecke_nur_ohne_ranking_bis_position = 20",
                 "serp_ueberschneidung = 50", "cluster_dichte = 50"):
        assert re.search(rf"^{re.escape(line)} #@param", source, flags=re.M), line
    assert re.search(r"^kalibrierung_bis_position = 5 #@param", _source("7a"), flags=re.M)


WEAK_QUERY = "getreidefreies trockenfutter hund"


def _verdict_of(nb, query):
    from qum import labels as L

    decisions = nb.ns["decisions"]
    return decisions.loc[decisions[L.C_QUERY] == query, L.C_VERDICT].item()


def _reason_of(nb, query):
    from qum import labels as L

    from qum.cannibal import cannibal_cases

    # die internen Fälle mit den Einstellungen, die Schritt 7c verwendet hat (das Blatt zeigt keine Gründe mehr)
    ns, settings = nb.ns, nb.ns["settings"]
    cannibal = cannibal_cases(ns["result"], ns["lead"], ns["threshold"], ns["rankings"],
                              settings["Rankt gut bis Position"], settings["Abstand fast gleich"])
    reasons = cannibal.loc[(cannibal[L.C_QUERY] == query) & (cannibal[L.C_STAGE] == L.STAGE_DANGER), L.C_REASON]
    assert reasons.nunique() == 1  # eine Zeile je URL, der Grund steht in jeder
    return reasons.iloc[0]


def test_smoke_step_7_margin_reaches_the_verdicts(nb):
    from qum import labels as L

    _load(nb)
    # die Kratzbaum-Seite rankt, erreicht die Schwelle, liegt aber 0,04 hinter der besten Seite
    weak = "Keyword;URL;Position\ngetreidefreies trockenfutter hund;https://www.tierbedarf.example/katzenzubehoer/kratzbaum;3\n"
    nb.uploads.append(("rankings.csv", weak.encode()))
    nb.run(5)
    _verdicts(nb, OWN, eigene_schwelle=0.2)
    assert _reason_of(nb, WEAK_QUERY) == L.REASON_BETTER
    _verdicts(nb, OWN, eigene_schwelle=0.2, abstand_fast_gleich=0.05)
    assert _reason_of(nb, WEAK_QUERY) == L.REASON_OK_CLOSE


def test_smoke_rerunning_only_step_7c_with_another_margin_changes_the_verdicts(nb, capsys):
    from qum import labels as L

    _load(nb)
    weak = "Keyword;URL;Position\ngetreidefreies trockenfutter hund;https://www.tierbedarf.example/katzenzubehoer/kratzbaum;3\n"
    nb.uploads.append(("rankings.csv", weak.encode()))
    nb.run(5)
    _verdicts(nb, OWN, eigene_schwelle=0.2)
    assert _reason_of(nb, WEAK_QUERY) == L.REASON_BETTER
    capsys.readouterr()
    nb.run("7c", abstand_fast_gleich=0.05)
    assert _reason_of(nb, WEAK_QUERY) == L.REASON_OK_CLOSE
    assert nb.ns["threshold"] == 0.2 and nb.ns["settings"]["Abstand fast gleich"] == 0.05
    assert "✅ Schritt 7c fertig. Weiter mit Schritt 8 (Export)." in capsys.readouterr().out
    nb.run(8)


def test_step_7b_threshold_dropdown_offers_the_three_choices():
    from qum import labels as L

    match = re.search(r'^schwelle_waehlen = (".*?") #@param (\[.*\])$', _source("7b"), flags=re.M)
    assert json.loads(match.group(1)) == CALIBRATED
    assert json.loads(match.group(2)) == CHOICES == list(L.THRESHOLD_CHOICE)
    assert [L.THRESHOLD_CHOICE[c] for c in CHOICES] == ["rankings", "median", "manuell"]
    assert _form_names("7a") == {"kalibrierung_bis_position"}
    assert _form_names("7b") == {"schwelle_waehlen", "eigene_schwelle"}


def _many_rankings():
    """24 Paare bis Position 2: genug für die Kalibrierung."""
    queries = (EXAMPLES / "queries.csv").read_text(encoding="utf-8").splitlines()[1:]
    urls = [line.split(";")[0] for line in (EXAMPLES / "frog_export.csv").read_text(encoding="utf-8").splitlines()[1:3]]
    lines = ["Keyword;URL;Position"] + [f"{q};{u};{n}" for q in queries for n, u in enumerate(urls, start=1)]
    return ("rankings.csv", ("\n".join(lines) + "\n").encode())


def test_smoke_step_7a_shows_proposals_and_builds_no_verdicts(nb, capsys):
    shown = []
    nb.ns["display"] = shown.append
    _load(nb)
    nb.uploads.append(_many_rankings())
    nb.run(5)
    _verdicts(nb)
    shown.clear()
    capsys.readouterr()
    nb.run("7a")
    out = capsys.readouterr().out
    assert "Die Schwelle ist die Cosinus-Ähnlichkeit zwischen Query und Seite, ab der eine Seite als passend gilt." in out
    assert re.search(r"Aus Rankings kalibriert: 0\.\d{4}\. Diesen Score erreichen 75 % der 24 Paare", out)
    assert re.search(r"Mittlerer bester Score \(nicht kalibriert\): 0\.\d{4}\.", out)
    assert "per Konstruktion darunter" in out
    assert out.rstrip().endswith("✅ Schritt 7a fertig. Weiter mit Schritt 7b: Schwelle festlegen.")
    assert len(shown) == 2  # Beispielpaare um beide Vorschläge
    for name in ("decisions", "cannibal", "settings", "threshold", "threshold_source", "threshold_label"):
        assert nb.ns[name] is None, name
    with pytest.raises(colab.NotebookStop):
        nb.run(8)
    assert NEED_7C in capsys.readouterr().out


@pytest.mark.parametrize(
    "choice, source, form",
    [(CALIBRATED, "rankings", {}), (MEDIAN, "median", {}), (OWN, "manuell", {"eigene_schwelle": 0.42})],
)
def test_smoke_step_7_each_choice_builds_verdicts(nb, capsys, choice, source, form):
    from qum.threshold import calibrated_threshold, median_threshold

    _load(nb)
    nb.uploads.append(_many_rankings())
    nb.run(5)
    _verdicts(nb, choice, **form)
    out = capsys.readouterr().out
    expected = {
        "rankings": lambda: calibrated_threshold(nb.ns["result"], nb.ns["lead"], nb.ns["rankings"]).value,
        "median": lambda: median_threshold(nb.ns["result"], nb.ns["lead"]).value,
        "manuell": lambda: 0.42,
    }[source]()
    assert nb.ns["decisions"] is not None and nb.ns["cannibal"] is not None
    assert nb.ns["threshold_source"] == source and nb.ns["threshold_label"] == choice
    assert nb.ns["threshold"] == expected
    assert nb.ns["settings"]["Schwelle aus"] == choice
    assert nb.ns["settings"]["Schwelle"] == expected
    assert f"Verwendete Schwelle {expected:.4f}" in out
    assert f"✅ Schritt 7b fertig: Schwelle {expected:.4f}. Weiter mit Schritt 7c: Urteile bilden." in out
    assert out.rstrip().endswith("✅ Schritt 7c fertig. Weiter mit Schritt 8 (Export).")


def test_smoke_calibration_with_too_few_pairs_stops(nb, capsys):
    _load(nb)
    nb.run("7a")
    capsys.readouterr()
    with pytest.raises(colab.NotebookStop):
        nb.run("7b")  # Vorgabe: Aus Rankings kalibriert
    assert "Es sind keine Rankings geladen (Schritt 5 oder 6)." in capsys.readouterr().out
    nb.run(5, "rankings.csv")
    nb.run("7a")
    capsys.readouterr()
    with pytest.raises(colab.NotebookStop):
        nb.run("7b", schwelle_waehlen=CALIBRATED)
    out = capsys.readouterr().out
    assert "❌ Die Kalibrierung aus Rankings ist nicht verfügbar" in out
    assert "Nur 4 Ranking-Paare bis Position 5, für eine Kalibrierung sind 20 nötig." in out
    assert "Wähle bei schwelle_waehlen eine andere Option und starte die Zelle erneut." in out
    assert nb.ns["threshold"] is None and nb.ns["decisions"] is None


@pytest.mark.parametrize("value", [0, -0.2, 1.5])
def test_smoke_invalid_own_threshold_stops(nb, capsys, value):
    _load(nb)
    nb.run("7a")
    with pytest.raises(colab.NotebookStop):
        nb.run("7b", schwelle_waehlen=OWN, eigene_schwelle=value)
    assert "❌ Die eigene Schwelle muss größer als 0 und höchstens 1 sein." in capsys.readouterr().out
    assert nb.ns.get("threshold") is None and nb.ns.get("decisions") is None


def test_smoke_own_threshold_of_one_is_accepted(nb):
    _load(nb)
    _verdicts(nb, OWN, eigene_schwelle=1.0)
    assert nb.ns["settings"]["Schwelle"] == 1.0


def test_install_line_pins_the_package_version():
    import tomllib

    import qum

    source = next(s for s in _code_cells() if s.startswith("#@title Schritt 2"))
    match = re.search(
        r'^!pip install -q "qum\[local\] @ git\+https://github\.com/DanielKremer91/query-url-matcher@v([^"]+)"$',
        source,
        flags=re.M,
    )
    assert match and match.group(1) == qum.__version__
    pyproject = tomllib.loads((ROOT / "pyproject.toml").read_text(encoding="utf-8"))
    assert pyproject["project"]["version"] == qum.__version__
    committed = (ROOT / "query_url_matcher.ipynb").read_text(encoding="utf-8")
    assert f"query-url-matcher@v{qum.__version__}" in committed



@pytest.mark.parametrize("value", [-0.01, 0.1, 2])
def test_smoke_margin_outside_its_range_stops(nb, capsys, value):
    _load(nb)
    _threshold(nb)
    with pytest.raises(colab.NotebookStop):
        nb.run("7c", abstand_fast_gleich=value)
    out = capsys.readouterr().out
    assert "Differenz von Cosinus-Scores (zum Beispiel 0.01), kein Prozentwert" in out
    assert KEPT not in out


def test_smoke_margin_zero_is_accepted(nb):
    _load(nb)
    _verdicts(nb, abstand_fast_gleich=0)
    assert nb.ns["settings"]["Abstand fast gleich"] == 0


def _closing_line(out):
    return next(line for line in out.splitlines() if line.startswith("ℹ️ Noch keine Urteile"))


def test_smoke_proposals_closing_line_lists_only_available_options(nb, capsys):
    _load(nb)
    nb.run("7a")
    line = _closing_line(capsys.readouterr().out)
    assert "Aus Rankings kalibriert" not in line and "Schritt 7b" in line
    assert '"Mittlerer bester Score (nicht kalibriert)"' in line and '"Eigener Wert"' in line
    nb.uploads.append(_many_rankings())
    nb.run(5)
    nb.run("7a")
    assert '"Aus Rankings kalibriert"' in _closing_line(capsys.readouterr().out)


def test_smoke_step_8_before_step_7c_points_to_step_7c_after_each_sub_step(nb, capsys):
    _load(nb)
    nb.run("7a")
    capsys.readouterr()
    with pytest.raises(colab.NotebookStop):
        nb.run(8)
    assert NEED_7C in capsys.readouterr().out
    nb.run("7b", schwelle_waehlen=MEDIAN)
    capsys.readouterr()
    with pytest.raises(colab.NotebookStop):
        nb.run(8)
    assert NEED_7C in capsys.readouterr().out
    nb.run(4)  # ein früherer Schritt setzt den Zustand zurück
    with pytest.raises(colab.NotebookStop):
        nb.run(8)
    assert NEED_7C in capsys.readouterr().out
    _verdicts(nb)
    nb.run(8)
    assert "✅ Schritt 8 fertig" in capsys.readouterr().out


def test_smoke_step_7b_before_step_7a_stops(nb, capsys):
    _load(nb)
    with pytest.raises(colab.NotebookStop):
        nb.run("7b", schwelle_waehlen=MEDIAN)
    assert NEED_7A in capsys.readouterr().out
    assert nb.ns.get("threshold") is None


def test_smoke_step_7c_before_step_7b_stops(nb, capsys):
    _load(nb)
    with pytest.raises(colab.NotebookStop):
        nb.run("7c")
    assert NEED_7B in capsys.readouterr().out
    nb.run("7a")
    with pytest.raises(colab.NotebookStop):
        nb.run("7c")
    assert NEED_7B in capsys.readouterr().out
    assert nb.ns.get("decisions") is None


def test_smoke_rerunning_step_7b_invalidates_the_verdicts(nb, capsys):
    _load(nb)
    _verdicts(nb)
    nb.run("7b", schwelle_waehlen=OWN, eigene_schwelle=0.3)
    assert nb.ns["threshold"] == 0.3 and nb.ns["threshold_source"] == "manuell"
    for name in ("decisions", "cannibal", "gaps", "new_pages", "settings"):
        assert nb.ns[name] is None, name
    capsys.readouterr()
    with pytest.raises(colab.NotebookStop):
        nb.run(8)
    assert NEED_7C in capsys.readouterr().out
    nb.run("7c")
    assert nb.ns["settings"]["Schwelle"] == 0.3 and nb.ns["settings"]["Schwelle aus"] == OWN
    nb.run(8)


@pytest.mark.parametrize("step", [5, 6])
def test_smoke_new_rankings_or_serps_invalidate_the_threshold(nb, capsys, step):
    _load(nb)
    _verdicts(nb)
    nb.run(step, {5: "rankings.csv", 6: "serps.csv"}[step])
    for name in ("calibration_position", "threshold", "threshold_source", "threshold_label", "decisions"):
        assert nb.ns[name] is None, name
    capsys.readouterr()
    for later, hint in (("7c", NEED_7B), ("7b", NEED_7A)):
        with pytest.raises(colab.NotebookStop):
            nb.run(later)
        assert hint in capsys.readouterr().out


@pytest.mark.parametrize(
    "step, form",
    [
        ("7a", {"kalibrierung_bis_position": 0}),
        ("7b", {"schwelle_waehlen": OWN, "eigene_schwelle": 0}),
        ("7b", {"schwelle_waehlen": CALIBRATED}),
        ("7c", {"abstand_fast_gleich": 0.5}),
        ("7c", {"rankt_gut_bis_position": 0}),
    ],
)
def test_smoke_rejected_step_7_says_the_last_verdicts_still_apply(nb, capsys, step, form):
    _load(nb)
    for earlier in ("7a", "7b", "7c")[: ("7a", "7b", "7c").index(step)]:
        nb.run(earlier, **({"schwelle_waehlen": MEDIAN} if earlier == "7b" else {}))
    capsys.readouterr()
    with pytest.raises(colab.NotebookStop):
        nb.run(step, **form)
    assert KEPT not in capsys.readouterr().out
    _verdicts(nb)
    before = {name: nb.ns[name] for name in ("decisions", "threshold", "calibration_position")}
    capsys.readouterr()
    with pytest.raises(colab.NotebookStop):
        nb.run(step, **form)
    out = capsys.readouterr().out
    assert out.startswith("❌ ") and out.rstrip().endswith(KEPT)
    assert all(nb.ns[name] is value for name, value in before.items())


def test_notebook_requests_a_gpu_runtime():
    meta = build()["metadata"]
    assert meta["accelerator"] == "GPU"
    assert meta["colab"]["gpuType"] == "T4"


def test_smoke_step_2_hides_hugging_face_token_notice(nb, monkeypatch, capsys):
    import logging

    http_logger = logging.getLogger("huggingface_hub.utils._http")
    monkeypatch.setattr(http_logger, "level", logging.NOTSET)
    _run_step_2(nb, monkeypatch, MODELS["e5-large"].label, gpu=True)
    http_logger.warning("Warning: You are sending unauthenticated requests to the HF Hub.")
    captured = capsys.readouterr()
    assert "unauthenticated" not in captured.out + captured.err
    assert http_logger.getEffectiveLevel() >= logging.ERROR


def test_upload_steps_explain_where_the_upload_button_appears():
    for source in _code_cells():
        if "colab.upload(" in source:
            first_hint = next(line for line in source.splitlines() if line.startswith("#@markdown"))
            assert "Play-Symbol" in first_hint and "Dateien auswählen" in first_hint, source.splitlines()[0]


def test_smoke_step_3_names_the_columns_actually_used(nb, capsys):
    nb.run(3, "queries.csv", "frog_export.csv")
    out = capsys.readouterr().out
    assert "URL = Address, Content = Extract Main Content 1" in out
    assert "None" not in out


def test_step_4_says_chunking_is_always_on():
    source = next(s for s in _code_cells() if s.startswith("#@title Schritt 4"))
    assert "Chunking ist immer aktiv" in source and "0 heißt: empfohlene Größe" in source


def test_step_4_explains_basis_options_and_has_no_top_n():
    source = next(s for s in _code_cells() if s.startswith("#@title Schritt 4"))
    for text in ("**Chunk** (Empfehlung)", "**Gesamt-URL:**", "**Kombi:**"):
        assert text in source, text
    assert "top_n" not in source and _form_names(4) == {"chunk_groesse", "chunk_overlap", "bewertungsgrundlage", "kombi_gewicht_chunk"}


@pytest.mark.parametrize("basis, column", [("Chunk", "Score Chunk"), ("Kombi", "Score Kombi")])
def test_smoke_step_4_previews_the_first_ten_queries_from_the_overview_columns(nb, capsys, basis, column):
    from qum import labels as L
    from qum.match import best_matches

    shown = []
    nb.ns["display"] = shown.append
    nb.run(3, "queries.csv", "frog_export.csv")
    nb.run(4, bewertungsgrundlage=basis)
    (frame,) = shown
    assert list(frame.columns) == [L.C_QUERY, L.C_BEST_URL, L.C_SCORE, L.C_SECOND_URL, L.C_SCORE_2]
    assert len(frame) == 10
    matches = best_matches(nb.ns["result"], nb.ns["lead"], nb.ns["weight"]).head(10)
    assert frame[L.C_SCORE].tolist() == matches[column].tolist()
    assert frame[L.C_SECOND_URL].tolist() == matches[L.C_SECOND_URL].tolist()
    assert "✅ Schritt 4 fertig: 12 Queries gematcht." in capsys.readouterr().out


def test_upload_steps_say_the_column_fields_normally_stay_empty():
    for source in _code_cells():
        if "colab.upload(" in source:
            assert "Die Felder unten bleiben normalerweise leer." in source, source.splitlines()[0]


def test_smoke_steps_5_and_6_name_the_detected_columns(nb, capsys):
    _load(nb)
    nb.run(5, "rankings.csv")
    nb.run(6, "serps.csv")
    out = capsys.readouterr().out
    assert "Erkannte Spalten: Keyword = Keyword, URL = URL, Position = Position" in out
    assert "Erkannte Spalten: Keyword = Keyword, URL = URL, Position = Position, Type = Type" in out


def test_step_7_explains_every_field():
    for step in ("7a", "7b", "7c"):
        source = _source(step)
        for name in _form_names(step):
            assert f"**{name}" in source, name
    assert "75 %" in _source("7a")


def _first_hint(step):
    return next(line for line in _source(step).splitlines() if line.startswith("#@markdown"))


def test_step_7_cells_say_how_to_start():
    assert _first_hint("7a") == "#@markdown **▶ Einfach starten, hier ist nichts einzutragen.**"
    assert _first_hint("7c") == (
        "#@markdown **▶ Starten. Die Werte darunter kannst du für den ersten Lauf auf den Voreinstellungen lassen.**"
    )


def test_smoke_step_7_explains_the_verdicts_it_shows(nb, capsys):
    _load(nb)
    _verdicts(nb)
    out = capsys.readouterr().out
    assert "Was die Urteile bedeuten:" in out
    assert "Content-Lücke: Keine Seite erreicht die Schwelle, und keine eigene Seite rankt" in out


def test_step_7_explains_the_threshold_and_every_option():
    assert "bildet noch keine Urteile" in _source("7a")
    source = _source("7b")
    for text in ("**Aus Rankings kalibriert**", "am verlässlichsten", "Schritt 5 oder 6",
                 "**Mittlerer bester Score (nicht kalibriert):**", "per Konstruktion darunter",
                 "**Eigener Wert:**", "Beispielen aus Schritt 7a"):
        assert text in source, text


def test_notebook_and_readme_name_the_sub_steps_of_step_7():
    texts = {"notebook": "\n".join(source for _, source in CELLS), "README": (ROOT / "README.md").read_text(encoding="utf-8")}
    for name, text in texts.items():
        assert not re.search(r"Schritt 7(?![abc])", text), name
        assert "schwelle_bestimmen" not in text, name
    intro = next(source for kind, source in CELLS if kind == "markdown")
    readme_start = texts["README"].split("## So startest du")[1].split("\n## ")[0]
    for text in (intro, readme_start):
        assert all(f"7{x}" in text for x in "abc")


def test_step_3_explains_both_input_files():
    source = next(s for s in _code_cells() if s.startswith("#@title Schritt 3"))
    for text in ("**Queries-Datei:**", "**Frog-Export:**", "Custom Extraction", "Configuration → Custom → Custom Extraction", "Address"):
        assert text in source, text


def test_step_7c_explains_both_cluster_sliders_with_examples():
    source = next(s for s in _code_cells() if s.startswith("#@title Schritt 7c"))
    for text in ("**serp_ueberschneidung:**", "6 von 10 URLs", "**cluster_dichte:**", "verhindert Ketten", "A und E fallen raus",
                 "ändern nur die Spalte Thema im Blatt „Potentielle Content-Lücken“"):
        assert text in source, text


def test_step_7c_explains_the_gap_fields():
    source = _source("7c")
    assert "luecke_unter_score" not in source and "eigene Einstellungen" not in source
    for text in ("**luecke_nur_ohne_ranking_bis_position:**",
                 "0 = aus", "Ohne Rankings wird der Wert ignoriert", "Rankt trotz schwachem Match"):
        assert text in source, text


def test_smoke_step_7c_gap_sheet_lists_exactly_the_content_gap_verdicts(nb, capsys):
    from qum import labels as L

    _load(nb)
    _verdicts(nb)
    out = capsys.readouterr().out
    gaps, decisions = nb.ns["gaps"], nb.ns["decisions"]
    assert set(gaps[L.C_QUERY]) == set(decisions.loc[decisions[L.C_VERDICT] == L.V_GAP, L.C_QUERY])
    assert f"{len(gaps):>5} potentielle Content-Lücken (Schwelle nicht kalibriert)" in out
    assert "neue Seiten" not in out  # ohne SERPs gibt es kein Thema
    assert "eigenen Einstellungen" not in out and "Lücke unter Score" not in nb.ns["settings"]
    assert gaps[L.C_TO_THRESHOLD].is_monotonic_increasing
    assert gaps[L.C_TOPIC].isna().all()


def test_smoke_step_7c_gaps_with_rankings_and_serps(nb, capsys):
    from qum import labels as L
    from qum.gaps import count_new_pages

    _load(nb)
    nb.run(5, "rankings.csv")
    nb.run(6, "serps.csv")
    _verdicts(nb, OWN, eigene_schwelle=1.0, luecke_nur_ohne_ranking_bis_position=0)
    out = capsys.readouterr().out
    gaps, decisions = nb.ns["gaps"], nb.ns["decisions"]
    # gut rankende Queries sind nie eine Lücke, alle anderen schon (Schwelle 1.0: keine Seite passt)
    good = set(nb.ns["rankings"].query("position <= 10")["query_norm"])
    expected = {q for q in nb.ns["queries"] if q.lower() not in good}
    assert set(gaps[L.C_QUERY]) == expected and gaps[L.C_TOPIC].notna().any()
    assert nb.ns["new_pages"] == count_new_pages(gaps) < len(gaps)
    assert f"   {len(gaps):>5} potentielle Content-Lücken, zusammen {nb.ns['new_pages']} neue Seiten" in out
    nb.run("7c", luecke_nur_ohne_ranking_bis_position=40)
    ranking = set(nb.ns["rankings"].query("position <= 40")["query_norm"])
    assert set(nb.ns["gaps"][L.C_QUERY]) == {q for q in nb.ns["queries"] if q.lower() not in ranking}
    assert nb.ns["settings"]["Lücke nur ohne Ranking bis Position"] == 40
    decisions = nb.ns["decisions"]
    assert set(nb.ns["gaps"][L.C_QUERY]) == set(decisions.loc[decisions[L.C_VERDICT] == L.V_GAP, L.C_QUERY])


def test_smoke_step_7c_ignores_the_ranking_condition_without_rankings(nb, capsys):
    _load(nb)
    _verdicts(nb, OWN, eigene_schwelle=1.0, luecke_nur_ohne_ranking_bis_position=5)
    out = capsys.readouterr().out
    assert "ℹ️ luecke_nur_ohne_ranking_bis_position wird ignoriert: Es sind keine Rankings geladen." in out
    assert len(nb.ns["gaps"]) == 12
    assert nb.ns["settings"]["Lücke nur ohne Ranking bis Position"] == "ignoriert, keine Rankings"


@pytest.mark.parametrize(
    "form, hint",
    [({"luecke_nur_ohne_ranking_bis_position": -1}, "luecke_nur_ohne_ranking_bis_position")],
)
def test_smoke_step_7c_rejects_invalid_gap_settings(nb, capsys, form, hint):
    _load(nb)
    _threshold(nb)
    capsys.readouterr()
    with pytest.raises(colab.NotebookStop):
        nb.run("7c", **form)
    out = capsys.readouterr().out
    assert out.startswith("❌ ") and hint in out


def test_verdict_guide_sits_before_step_7a_and_covers_every_verdict():
    from qum import labels as L

    kinds = [(kind, source.splitlines()[0]) for kind, source in CELLS]
    guide_index = next(i for i, (kind, first) in enumerate(kinds) if kind == "markdown" and "So entstehen die Urteile" in first)
    assert kinds[guide_index + 1][1].startswith("#@title Schritt 7a")
    guide = CELLS[guide_index][1]
    for verdict in (L.V_OK, L.V_CANNIBAL, L.V_WATCH, L.V_GAP, L.V_MATCH):
        assert verdict in guide, verdict
    assert "| **Rankt gut**" in guide


def test_smoke_step_7c_counts_the_queries_of_the_cannibalisation_sheet(nb, capsys):
    from qum import labels as L

    _load(nb)
    nb.run(5, "rankings.csv")
    nb.run("7a")
    nb.run("7b", schwelle_waehlen=MEDIAN)
    nb.run("7c")
    decisions, cannibal = nb.ns["decisions"], nb.ns["cannibal"]
    assert L.C_PRIORITY not in decisions.columns
    assert set(decisions.loc[decisions[L.C_VERDICT] == L.V_CANNIBAL, L.C_QUERY]) <= set(cannibal[L.C_QUERY])
    out = capsys.readouterr().out
    assert f"{cannibal[L.C_QUERY].nunique():>5} Queries im Blatt Kannibalisierungsgefahr, über alle Urteile (" in out


def test_verdict_guide_explains_the_cannibalisation_sheet_and_its_urgency():
    from qum import labels as L

    guide = next(source for kind, source in CELLS if kind == "markdown" and "So entstehen die Urteile" in source)
    assert "Blatt „Kannibalisierungsgefahr“ (unabhängig vom Urteil)" in guide
    for value in (L.PRIO_VERY_HIGH, L.PRIO_HIGH, L.PRIO_MID, L.PRIO_LOW, L.PRIO_VERY_LOW, L.PRIO_OPEN):
        assert f"| **{value}** |" in guide, value
    assert "**ja**" not in guide and "**möglich**" not in guide and "Chunk auf anderer Seite" in guide


def test_notebook_ends_with_step_8_and_nothing_mentions_step_9_or_pairs():
    assert CELLS[-1][1].startswith("#@title Schritt 8")
    texts = {
        "notebook": "\n".join(source for _, source in CELLS),
        "README": (ROOT / "README.md").read_text(encoding="utf-8"),
        "spec": (ROOT / "docs/superpowers/specs/2026-10-05-query-url-matcher-design.md").read_text(encoding="utf-8"),
    }
    for name, text in texts.items():
        assert "Schritt 9" not in text and "Paare bewerten" not in text and "Keyword-URL-Paare" not in text, name


def test_no_serp_neighbour_hint_and_no_advice_column():
    from qum import labels as L

    for name in ("V_CHECK", "C_CAND", "C_CAND_KW", "C_CAND_OVERLAP", "C_CAND_SCORE", "C_CAND_CHUNK", "C_CAND_POS",
                 "C_CAND_MORE", "C_ADVICE"):
        assert not hasattr(L, name), name
    text = "\n".join(source for _, source in CELLS)
    assert "Vor Neuerstellung prüfen" not in text and "Die Empfehlung je Query" not in text


def _guide():
    return next(source for kind, source in CELLS if kind == "markdown" and "So entstehen die Urteile" in source)


def test_verdict_guide_shows_two_scenarios_without_and_with_rankings():
    from qum import labels as L

    guide = _guide()
    without = guide.index("### Ohne Rankings")
    with_rankings = guide.index("### Mit Rankings")
    column = guide.index("### Blatt „Kannibalisierungsgefahr“")
    assert without < with_rankings < column
    for verdict in (L.V_MATCH, L.V_CANNIBAL, L.V_GAP):
        assert verdict in guide[without:with_rankings], verdict
    for verdict in (L.V_OK, L.V_WATCH):
        assert verdict in guide[with_rankings:column], verdict
    assert L.V_WATCH not in guide[without:with_rankings]
    assert "„In Ordnung“ gibt es deshalb nur mit Rankings" in guide[without:with_rankings]
    assert "wie ohne Rankings" in guide[with_rankings:column]
    assert "die rankende Seite passt, aber eine andere passt besser" in guide[with_rankings:column]
    assert "nicht im Frog-Export" in guide[with_rankings:column]
    assert "Bestehende Seite nutzen" not in guide


def test_step_6_says_serps_only_group_the_gaps_into_topics():
    source = _source(6)
    assert "nur genutzt" in source and "Spalte Thema" in source and "Zahl neuer Seiten" in source and "Potentielle Content-Lücken" in source


def test_step_8_describes_the_four_sheets_one_line_each():
    from qum import export

    hints = [line for line in _source(8).splitlines() if line.startswith("#@markdown")]
    for name in (export.SHEET_OVERVIEW, export.SHEET_CANNIBAL, export.SHEET_GAPS, export.SHEET_README):
        assert sum(f"**{name}:**" in line for line in hints) == 1, name


def test_intro_and_readme_describe_the_four_sheets():
    from qum import export

    intro = next(source for kind, source in CELLS if kind == "markdown")
    readme = (ROOT / "README.md").read_text(encoding="utf-8")
    result = readme.split("## Ergebnis")[1].split("\n## ")[0]
    for text in (intro.split("## Ergebnis")[1].split("\n## ")[0], result):
        for name in (export.SHEET_OVERVIEW, export.SHEET_CANNIBAL, export.SHEET_GAPS, export.SHEET_README):
            assert name in text, name
        assert "Entscheidung" not in text and "Lücken je Cluster" not in text


def test_step_4_marks_the_weight_slider_as_only_relevant_for_kombi():
    source = _source(4)
    assert "**kombi_gewicht_chunk (nur relevant bei Kombi):**" in source
    assert source.index("(nur relevant bei Kombi)") < source.index("kombi_gewicht_chunk = 0.7")
    assert "nicht das Urteil" in source


def test_smoke_step_4_hints_when_the_weight_is_changed_but_not_used(nb, capsys):
    nb.run(3, "queries.csv", "frog_export.csv")
    hint = "wirkt nur bei Kombi"
    capsys.readouterr()
    nb.run(4, bewertungsgrundlage="Chunk", kombi_gewicht_chunk=0.5)
    assert hint in capsys.readouterr().out
    nb.run(4, bewertungsgrundlage="Gesamt-URL", kombi_gewicht_chunk=0.9)
    assert hint in capsys.readouterr().out
    nb.run(4, bewertungsgrundlage="Kombi", kombi_gewicht_chunk=0.5)
    assert hint not in capsys.readouterr().out
    nb.run(4, bewertungsgrundlage="Chunk")
    assert hint not in capsys.readouterr().out


# --- Versionsprüfung: Notebook und geladenes Paket müssen zusammenpassen -------------------------------------------


def test_every_code_cell_checks_the_loaded_package_version_before_using_it():
    from qum import __version__

    check = f'if _qum_paket.__version__ != "{__version__}":'
    for source in _code_cells():
        lines = source.splitlines()
        assert check in lines, lines[0]
        first_use = next(i for i, line in enumerate(lines) if line.startswith(("from qum", "import qum")))
        assert lines[first_use] == "import qum as _qum_paket", lines[0]
    step_2 = _source(2)
    assert step_2.index("!pip install") < step_2.index(check)


def test_smoke_a_stale_package_stops_with_a_restart_hint(nb, monkeypatch, capsys):
    import qum
    from qum import colab

    monkeypatch.setattr(qum, "__version__", "0.0.1")
    with pytest.raises(colab.NotebookStop):
        nb.run(3, "queries.csv", "frog_export.csv")
    out = capsys.readouterr().out
    assert "geladen ist aber 0.0.1" in out and "Laufzeit → Sitzung neu starten" in out and "ab Schritt 2" in out


def test_smoke_step_8_colours_the_lead_with_the_margin_of_step_7c(nb, tmp_path):
    from openpyxl import load_workbook

    from qum import labels as L

    _load(nb)
    _verdicts(nb, abstand_fast_gleich=0.05)
    nb.run(8)
    sheet = load_workbook(tmp_path / "query_url_matcher.xlsx")["Übersicht"]
    header = [cell.value for cell in sheet[1]]
    column = header.index(L.C_LEAD_GAP) + 1
    for row in range(2, sheet.max_row + 1):
        value = sheet.cell(row=row, column=column).value
        colour = sheet.cell(row=row, column=column).fill.fgColor.rgb[-6:]
        if value is not None and value <= 0.05:
            assert colour == "F8CBAD", (row, value)
