import io
import zipfile

import pandas as pd
import pytest
from openpyxl import load_workbook

from qum import export
from qum import labels as L


def _decisions(with_cluster=False):
    df = pd.DataFrame(
        {
            L.C_QUERY: ["a", "b", "c"],
            L.C_VERDICT: [L.V_MATCH, L.V_GAP, L.V_CHECK],
            L.C_BEST_URL: ["u1", "u2", "u3"],
        }
    )
    if with_cluster:
        df[L.C_CLUSTER] = [1, 1, 0]
    return df


TOP = pd.DataFrame({L.C_QUERY: ["a"], L.C_URL: ["u1"]})
CANNIBAL = pd.DataFrame(columns=[L.C_QUERY, L.C_STAGE, L.C_REASON, L.C_COMPETING])
SETTINGS = {"Modell": "multilingual-e5-large", "Schwelle": "0.81"}


def test_content_gaps_contains_gap_and_check():
    assert export.content_gaps(_decisions())[L.C_QUERY].tolist() == ["b", "c"]


def test_sheets_without_serps():
    sheets = export.build_sheets(_decisions(), TOP, CANNIBAL, SETTINGS)
    assert list(sheets) == [
        export.SHEET_README, export.SHEET_DECISION, export.SHEET_TOP, export.SHEET_CANNIBAL, export.SHEET_GAPS,
    ]


def test_sheets_with_serps():
    sheets = export.build_sheets(_decisions(with_cluster=True), TOP, CANNIBAL, SETTINGS)
    assert list(sheets)[-1] == export.SHEET_GAP_CLUSTERS
    assert not hasattr(export, "SHEET_PAIRS")


def test_readme_lists_present_sheets_settings_and_disclaimer():
    readme = export.build_sheets(_decisions(), TOP, CANNIBAL, SETTINGS)[export.SHEET_README]
    assert list(readme.columns) == [L.R_AREA, L.R_ENTRY, L.R_TEXT]
    assert [L.R_AREA, L.R_ENTRY, L.R_TEXT] == ["Bereich", "Eintrag", "Erklärung"]
    entries = readme[L.R_ENTRY].tolist()
    assert export.SHEET_GAPS in entries
    assert "Modell" in entries
    assert L.V_CANNIBAL in entries
    assert readme.iloc[0][L.R_TEXT].startswith("Das Notebook sortiert vor")


def test_every_sheet_has_query_first_except_readme_and_summary():
    sheets = export.build_sheets(_decisions(with_cluster=True), TOP, CANNIBAL, SETTINGS)
    for name, df in sheets.items():
        if name not in (export.SHEET_README, export.SHEET_GAP_CLUSTERS):
            assert df.columns[0] == L.C_QUERY, name


def test_every_label_column_has_a_help_text():
    # C_SIDE und C_SCORE gehören nur zu den Prüfbeispielen in Schritt 7, nicht zum Export
    columns = [v for k, v in vars(L).items() if k.startswith("C_") and k not in ("C_SIDE", "C_SCORE")]
    assert columns
    for column in columns:
        assert column in export._COLUMN_HELP, column


def test_readme_explains_candidate_columns_only_when_present():
    decisions = _decisions(with_cluster=True)
    decisions[L.C_CAND] = ["", "", "u1"]
    readme = export.build_sheets(decisions, TOP, CANNIBAL, SETTINGS)[export.SHEET_README]
    assert L.C_CAND in readme[L.R_ENTRY].tolist()
    stages = readme[readme[L.R_AREA] == "Stufe"][L.R_ENTRY].tolist()
    assert stages == [L.STAGE_DANGER, L.STAGE_VISIBLE]
    plain = export.build_sheets(_decisions(), TOP, CANNIBAL, SETTINGS)[export.SHEET_README]
    assert L.C_CAND not in plain[L.R_ENTRY].tolist()
    assert L.C_CLUSTER not in plain[L.R_ENTRY].tolist()


def test_write_excel_creates_sheets_and_colours_verdicts(tmp_path):
    path = tmp_path / "out.xlsx"
    export.write_excel(path, export.build_sheets(_decisions(), TOP, CANNIBAL, SETTINGS))
    book = load_workbook(path)
    assert book.sheetnames[:2] == [export.SHEET_README, export.SHEET_DECISION]
    sheet = book[export.SHEET_DECISION]
    assert sheet["A1"].font.bold
    fills = [sheet.cell(row=r, column=2).fill.fgColor.rgb for r in (2, 3, 4)]
    assert [f[-6:] for f in fills] == ["C6EFCE", "FFC7CE", "FFEB9C"]


def test_write_csv_zip_has_one_file_per_sheet(tmp_path):
    path = tmp_path / "out.zip"
    export.write_csv_zip(path, export.build_sheets(_decisions(), TOP, CANNIBAL, SETTINGS))
    with zipfile.ZipFile(path) as archive:
        assert archive.namelist() == [
            "lesehilfe.csv", "entscheidung.csv", "top_treffer.csv", "kannibalisierungsgefahr.csv", "content_luecken.csv",
        ]
        df = pd.read_csv(io.BytesIO(archive.read("entscheidung.csv")), encoding="utf-8-sig", sep=";")
    assert df[L.C_QUERY].tolist() == ["a", "b", "c"]


@pytest.mark.parametrize("sep", [";", ","])
def test_write_csv_zip_round_trips_with_either_separator(tmp_path, sep):
    path = tmp_path / "out.zip"
    sheets = export.build_sheets(_decisions(), TOP, CANNIBAL, SETTINGS)
    sheets[export.SHEET_DECISION].loc[0, L.C_QUERY] = "futter; nass, 10 kg"  # beide Trennzeichen im Text
    export.write_csv_zip(path, sheets, sep=sep)
    with zipfile.ZipFile(path) as archive:
        raw = archive.read("entscheidung.csv").decode("utf-8-sig")
        df = pd.read_csv(io.StringIO(raw), sep=sep)
    assert raw.splitlines()[0].count(sep) == len(sheets[export.SHEET_DECISION].columns) - 1
    assert list(df.columns) == list(sheets[export.SHEET_DECISION].columns)
    assert df[L.C_QUERY].tolist() == ["futter; nass, 10 kg", "b", "c"]


def test_write_csv_zip_defaults_to_semicolon(tmp_path):
    path = tmp_path / "out.zip"
    export.write_csv_zip(path, export.build_sheets(_decisions(), TOP, CANNIBAL, SETTINGS))
    with zipfile.ZipFile(path) as archive:
        header = archive.read("entscheidung.csv").decode("utf-8-sig").splitlines()[0]
    assert ";" in header


def test_write_excel_strips_control_characters_without_mutating_input(tmp_path):
    decisions = _decisions()
    decisions[L.C_CHUNK] = ["ok", "tab\x0bvertical", "unit\x1fsep"]
    path = tmp_path / "out.xlsx"
    export.write_excel(path, export.build_sheets(decisions, TOP, CANNIBAL, SETTINGS))
    sheet = load_workbook(path)[export.SHEET_DECISION]
    values = [sheet.cell(row=r, column=4).value for r in (2, 3, 4)]
    assert values == ["ok", "tab vertical", "unit sep"]
    assert decisions[L.C_CHUNK].tolist() == ["ok", "tab\x0bvertical", "unit\x1fsep"]


def test_candidate_position_help_refers_to_the_gap_query():
    text = export._COLUMN_HELP[L.C_CAND_POS]
    assert "Lücken-Query" in text
    assert "Nachbar-Keyword" not in text


def test_help_texts_do_not_overclaim_serp_similarity():
    texts = list(export._COLUMN_HELP.values()) + list(export.VERDICT_HELP.values())
    assert all("fast gleicher SERP" not in text for text in texts)
    assert "stark überlappender SERP" in export.VERDICT_HELP[L.V_CHECK]


def _notes(readme):
    return readme[readme[L.R_AREA] == "Hinweis"][L.R_TEXT].tolist()


def test_readme_always_explains_the_score_band():
    readme = export.build_sheets(_decisions(), TOP, CANNIBAL, SETTINGS)[export.SHEET_README]
    assert _notes(readme)[0].startswith("Das Notebook sortiert vor")
    assert export.CAVEAT_SCORES in _notes(readme)
    assert export.CAVEAT_SCORES.startswith("Scores eines Modells liegen in einem engen Band (bei e5 etwa 0,7 bis 0,9).")


def test_readme_flags_uncalibrated_median_threshold():
    readme = export.build_sheets(_decisions(), TOP, CANNIBAL, SETTINGS, threshold_source="median")[export.SHEET_README]
    notes = _notes(readme)
    assert notes == [export.DISCLAIMER, export.CAVEAT_SCORES, export.CAVEAT_THRESHOLD["median"]]
    assert notes[2].startswith("Schwelle nicht kalibriert:")


def test_readme_explains_threshold_calibrated_from_rankings():
    readme = export.build_sheets(_decisions(), TOP, CANNIBAL, SETTINGS, threshold_source="rankings")[export.SHEET_README]
    notes = _notes(readme)
    assert notes[-1] == export.CAVEAT_THRESHOLD["rankings"]
    assert "75 % der gut rankenden Paare" in notes[-1] and "Rankt trotz schwachem Match" in notes[-1]


def test_readme_says_when_threshold_was_set_by_hand():
    readme = export.build_sheets(_decisions(), TOP, CANNIBAL, SETTINGS, threshold_source="manuell")[export.SHEET_README]
    assert _notes(readme)[-1] == export.CAVEAT_THRESHOLD["manuell"]
    assert "von Hand" in export.CAVEAT_THRESHOLD["manuell"]


def test_write_excel_keeps_formula_like_text_as_text(tmp_path):
    decisions = _decisions()
    decisions[L.C_QUERY] = ["=cmd|x", "+49 hotline", "@home"]
    decisions[L.C_CHUNK] = ["= 5 Euro", "-20 % Rabatt", "normal"]
    path = tmp_path / "out.xlsx"
    export.write_excel(path, export.build_sheets(decisions, TOP, CANNIBAL, SETTINGS))
    sheet = load_workbook(path)[export.SHEET_DECISION]
    cells = [sheet.cell(row=r, column=c) for c in (1, 4) for r in (2, 3, 4)]
    assert [cell.value for cell in cells] == ["=cmd|x", "+49 hotline", "@home", "= 5 Euro", "-20 % Rabatt", "normal"]
    assert all(cell.data_type == "s" for cell in cells)


def test_help_texts_describe_the_margin_rule():
    assert "fast gleich" in export.VERDICT_HELP[L.V_OK]
    assert "deutlich besser" in export.VERDICT_HELP[L.V_CANNIBAL]
    assert "rankende Seite" in export._COLUMN_HELP[L.C_BEST_URL]
    assert "deutlich besser" in export._STAGE_HELP[L.STAGE_DANGER]


def test_advice_column_help_calls_it_a_hint_to_check():
    assert export._COLUMN_HELP[L.C_ADVICE] == "Fester Hinweis, was zu prüfen ist, kein generierter Text. Entscheiden muss ein Mensch."


@pytest.mark.parametrize("sep, number", [(";", "0,8123"), (",", "0.8123")])
def test_write_csv_zip_writes_decimal_comma_only_with_semicolon(tmp_path, sep, number):
    path = tmp_path / "out.zip"
    sheets = export.build_sheets(_decisions(), TOP, CANNIBAL, SETTINGS)
    sheets[export.SHEET_DECISION][L.C_S_CHUNK] = [0.8123, 0.5, 0.25]
    export.write_csv_zip(path, sheets, sep=sep)
    with zipfile.ZipFile(path) as archive:
        raw = archive.read("entscheidung.csv").decode("utf-8-sig")
    assert raw.splitlines()[1].endswith(f"{sep}{number}")
    df = pd.read_csv(io.StringIO(raw), sep=sep, decimal="," if sep == ";" else ".")
    assert df[L.C_S_CHUNK].tolist() == [0.8123, 0.5, 0.25]



def test_risk_help_covers_ranking_url_outside_the_export():
    assert "oder die rankende URL steht nicht im Frog-Export" in export.VERDICT_HELP[L.V_CANNIBAL]


def _text_number_sheets():
    decisions = _decisions()
    decisions[L.C_POSITION] = ["4.3", "12", ""]
    decisions[L.C_CAND_POS] = ["", "7.5", ""]
    decisions[L.C_ADVICE] = ["https://a.de/x.html rankt auf Position 4.3, semantisch passt https://a.de/y besser.", "", ""]
    cannibal = pd.DataFrame(
        {
            L.C_QUERY: ["a"],
            L.C_STAGE: [L.STAGE_DANGER],
            L.C_REASON: [L.REASON_BETTER],
            L.C_COMPETING: ["https://a.de/x.html (Position 4.3, Score 0.842) | https://a.de/y (Score -0.0123)"],
        }
    )
    settings = {"Datum": "2026-10-05", "Modell": "multilingual-e5-large", "Schwelle": 0.8123, "Treffer je Query (Top-N)": 5}
    return export.build_sheets(decisions, TOP, cannibal, settings)


def _csv_frames(path, sep):
    with zipfile.ZipFile(path) as archive:
        return {
            name: pd.read_csv(io.BytesIO(archive.read(name)), encoding="utf-8-sig", sep=sep, dtype=str, keep_default_na=False)
            for name in ("entscheidung.csv", "kannibalisierungsgefahr.csv", "lesehilfe.csv")
        }


def test_semicolon_csv_uses_decimal_comma_in_positions_settings_and_competing_urls(tmp_path):
    sheets = _text_number_sheets()
    export.write_csv_zip(tmp_path / "out.zip", sheets, sep=";")
    frames = _csv_frames(tmp_path / "out.zip", ";")
    decisions = frames["entscheidung.csv"]
    assert decisions[L.C_POSITION].tolist() == ["4,3", "12", ""]
    assert decisions[L.C_CAND_POS].tolist() == ["", "7,5", ""]
    assert decisions[L.C_ADVICE].tolist()[0] == "https://a.de/x.html rankt auf Position 4,3, semantisch passt https://a.de/y besser."
    assert frames["kannibalisierungsgefahr.csv"].iloc[0][L.C_COMPETING] == (
        "https://a.de/x.html (Position 4,3, Score 0,842) | https://a.de/y (Score -0,0123)"
    )
    readme = frames["lesehilfe.csv"].set_index(L.R_ENTRY)[L.R_TEXT]
    assert readme["Schwelle"] == "0,8123"
    assert readme["Datum"] == "2026-10-05" and readme["Modell"] == "multilingual-e5-large"
    assert readme["Treffer je Query (Top-N)"] == "5"
    assert sheets[export.SHEET_DECISION][L.C_POSITION].tolist() == ["4.3", "12", ""]  # Original unverändert


def test_comma_csv_and_excel_keep_the_decimal_point_in_text(tmp_path):
    sheets = _text_number_sheets()
    export.write_csv_zip(tmp_path / "out.zip", sheets, sep=",")
    frames = _csv_frames(tmp_path / "out.zip", ",")
    assert frames["entscheidung.csv"][L.C_POSITION].tolist() == ["4.3", "12", ""]
    assert "Position 4.3, Score 0.842" in frames["kannibalisierungsgefahr.csv"].iloc[0][L.C_COMPETING]
    assert "Position 4.3," in frames["entscheidung.csv"][L.C_ADVICE].tolist()[0]
    export.write_excel(tmp_path / "out.xlsx", sheets)
    book = load_workbook(tmp_path / "out.xlsx")
    position_col = list(sheets[export.SHEET_DECISION].columns).index(L.C_POSITION) + 1
    assert book[export.SHEET_DECISION].cell(row=2, column=position_col).value == "4.3"
