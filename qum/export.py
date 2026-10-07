import re
import zipfile

import pandas as pd
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from . import labels as L
from .cannibal import COLUMNS as CANNIBAL_COLUMNS
from .elsewhere import COLUMNS as ELSEWHERE_COLUMNS
from .gaps import COLUMNS as GAP_COLUMNS
from .verdict import OVERVIEW_COLUMNS

SHEET_OVERVIEW = "Übersicht"
SHEET_CANNIBAL = "Kannibalisierungsgefahr"
SHEET_GAPS = "Potentielle Content-Lücken"
SHEET_ELSEWHERE = "Chunk auf anderer Seite"
SHEET_README = "Lesehilfe"

_CSV_NAMES = {
    SHEET_OVERVIEW: "uebersicht.csv",
    SHEET_CANNIBAL: "kannibalisierungsgefahr.csv",
    SHEET_GAPS: "potentielle_content_luecken.csv",
    SHEET_ELSEWHERE: "chunk_auf_anderer_seite.csv",
    SHEET_README: "lesehilfe.csv",
}

_SHEET_HELP = {
    SHEET_OVERVIEW: (
        "Eine Zeile je Query: die drei am besten passenden URLs mit ihren Scores, die beste eigene Rankingposition "
        "und das Urteil."
    ),
    SHEET_CANNIBAL: (
        "Queries, bei denen weitere eigene Seiten passen oder mehrere eigene URLs ranken, unabhängig vom Urteil, mit einer "
        "Zeile je konkurrierender URL (die Query steht in jeder Zeile) und der Einordnung, wie dringend der Fall ist. "
        f"Enthält jede Query mit dem Urteil '{L.V_CANNIBAL}'."
    ),
    SHEET_GAPS: (
        f"Genau die Queries mit dem Urteil '{L.V_GAP}', sortiert nach Abstand zur Schwelle: die sichersten Lücken "
        "zuerst. Wer nur deutliche Lücken will, filtert die Spalte Abstand zur Schwelle (zum Beispiel kleiner als -0.03)."
    ),
    SHEET_ELSEWHERE: (
        "Queries, deren bester Chunk auf einer anderen Seite liegt als die beste Seite insgesamt (höchster Gesamt-URL-"
        "Score), unabhängig von Schwelle und Bewertungsgrundlage. Eine Seite dreht sich insgesamt um das Thema, die "
        "konkreteste Antwort steht aber auf einer anderen: Kandidaten für Kannibalisierung durch einzelne Abschnitte. "
        "Stärkster Chunk zuerst."
    ),
    SHEET_README: "Diese Erklärungen: Hinweise, Blätter, Urteile, Spalten und die Einstellungen des Laufs.",
}

_COLUMN_HELP = {
    # Übersicht
    L.C_QUERY: "Die Suchanfrage, um die es in der Zeile geht.",
    L.C_BEST_URL: "Die Seite mit dem höchsten Score für die Query, nach der gewählten Bewertungsgrundlage.",
    L.C_CHUNK: "Der Textblock der besten URL, der zur Query am besten passt.",
    # Chunk auf anderer Seite
    L.C_CHUNK_URL: "Die Seite, auf der der zur Query am besten passende Textblock steht (höchster Chunk-Score).",
    L.C_BEST_CHUNK: "Dieser Textblock.",
    L.C_S_BEST_CHUNK: "Cosinus-Ähnlichkeit zwischen Query und diesem Textblock.",
    L.C_S_FULL_CHUNK_URL: "Gesamt-URL-Score der Seite mit dem besten Chunk: wie gut sie als Ganzes passt.",
    L.C_OVERALL_URL: "Die Seite, die als Ganzes am besten passt (höchster Gesamt-URL-Score).",
    L.C_S_FULL_OVERALL: "Gesamt-URL-Score dieser Seite.",
    L.C_RANK_IS: (
        f"Welche Seite für die Query rankt: '{L.RANK_IS_CHUNK}', '{L.RANK_IS_OVERALL}', '{L.RANK_IS_OTHER}', "
        f"'{L.CMP_NOT_RANKING}' (keine eigene URL rankt) oder '{L.CMP_NOT_IN_EXPORT}'. Leer ohne Rankings."
    ),
    L.C_S_CHUNK: "Cosinus-Ähnlichkeit zwischen Query und dem am besten passenden Textblock der besten URL.",
    L.C_S_FULL: "Cosinus-Ähnlichkeit zwischen Query und dem gesamten Main Content der besten URL.",
    L.C_S_COMBI: "Gewichtete Mischung aus Chunk-Score und Gesamt-URL-Score der besten URL.",
    L.C_TO_THRESHOLD: (
        "Score der besten URL minus Schwelle aus Schritt 7b. Positiv: die Seite passt, negativ: wie weit sie davon "
        "entfernt ist. Lesbarer als der Score selbst, weil die Scores je nach Modell in ganz anderen Bereichen liegen "
        "(bei e5 fast immer zwischen 0.75 und 0.9)."
    ),
    L.C_LEAD_GAP: (
        "Score der besten URL minus Score der zweitbesten URL, nach der gewählten Bewertungsgrundlage. "
        "Ein großer Vorsprung heißt: eine Seite sticht klar heraus, auch wenn der absolute Score niedrig ist. "
        "Farbe im Excel, gemessen am Abstand 'fast gleich' aus Schritt 7c: rot = höchstens dieser Abstand (die "
        "zweitbeste URL ist praktisch gleichauf), gelb = bis zum Doppelten, grün = die beste URL liegt klar vorn. "
        "Leer, wenn es nur eine URL gibt."
    ),
    L.C_SECOND_URL: "Die Seite mit dem zweithöchsten Score für die Query, nach der gewählten Bewertungsgrundlage.",
    L.C_S_CHUNK_2: "Chunk-Score der zweitbesten URL.",
    L.C_S_FULL_2: "Gesamt-URL-Score der zweitbesten URL.",
    L.C_S_COMBI_2: "Kombi-Score der zweitbesten URL.",
    L.C_THIRD_URL: "Die Seite mit dem dritthöchsten Score für die Query, nach der gewählten Bewertungsgrundlage.",
    L.C_S_CHUNK_3: "Chunk-Score der drittbesten URL.",
    L.C_S_FULL_3: "Gesamt-URL-Score der drittbesten URL.",
    L.C_S_COMBI_3: "Kombi-Score der drittbesten URL.",
    L.C_POSITION: "Beste eigene Position für die Query laut Rankings. Leer, wenn die Query nicht rankt oder keine Rankings geladen sind.",
    L.C_RANK_URL: "Die eigene URL, die laut Rankings am besten für die Query rankt, egal auf welcher Position.",
    L.C_RANK_IS_BEST: (
        f"'{L.YES}': die rankende URL ist die beste URL. '{L.NO}': eine andere Seite hat einen höheren Score, auch wenn "
        "sie nur knapp vorn liegt; wie knapp, zeigen die Scores und der Grund im Blatt Kannibalisierungsgefahr. "
        f"'{L.CMP_NOT_RANKING}': keine eigene URL rankt für die Query. '{L.CMP_NOT_IN_EXPORT}': die rankende URL "
        "steht nicht im Frog-Export und wurde nicht verglichen. Leer ohne Rankings."
    ),
    L.C_VERDICT: "Einordnung der Query, siehe die Urteile weiter oben in dieser Lesehilfe.",
    # Kannibalisierungsgefahr
    L.C_PRIORITY: (
        f"Wie dringend der Fall ist, je Query (steht in jeder Zeile). Das Blatt ist danach sortiert, die dringendsten "
        f"Fälle zuerst, bei gleicher Einordnung die engste Konkurrenz zuerst. Konkurrieren Seiten eng (höchstens der Abstand "
        f"'fast gleich' hinter der besten, oder eine andere passende Seite ist besser als die rankende) oder ranken mehrere "
        f"eigene URLs bis sichtbar_bis_position, zählt das beste eigene Ranking: '{L.PRIO_VERY_HIGH}' knapp hinter den "
        f"Top-Rankings bis sichtbar_bis_position (Voreinstellung 20), fast oben, die Konkurrenz bremst vermutlich; "
        f"'{L.PRIO_HIGH}' schlechter oder kein Ranking, abgrenzen, zusammenführen oder Hauptseite festlegen; "
        f"'{L.PRIO_MID}' Top-Ranking (bis rankt_gut_bis_position) einer anderen als der besten Seite, prüfen, ob Google "
        f"die richtige gewählt hat; '{L.PRIO_LOW}' Top-Ranking der besten Seite, beobachten. Liegen die weiteren Seiten "
        f"deutlich dahinter: '{L.PRIO_MID}', wenn eine andere als die beste Seite rankt, sonst '{L.PRIO_VERY_LOW}'. "
        f"'{L.PRIO_OPEN}': keine Rankings geladen, nur semantisch geprüft."
    ),
    L.C_NO: "Laufende Nummer der URL innerhalb der Query. Jede konkurrierende URL hat eine eigene Zeile.",
    L.C_COMP_URL: (
        "Eine der konkurrierenden eigenen URLs: Seiten, die die Schwelle erreichen und nah an der besten liegen oder "
        "weiter dahinter, und eigene URLs, die bis sichtbar_bis_position ranken. Jede Seite einmal, nach Score sortiert, "
        "Seiten ohne Score (nicht im Frog-Export) zuletzt nach Position. Nach dieser Spalte filtern zeigt, bei welchen "
        "Queries eine Seite mit anderen konkurriert."
    ),
    L.C_COMP_SCORE: "Score dieser URL nach der gewählten Bewertungsgrundlage. Leer, wenn sie nicht im Frog-Export steht.",
    L.C_GAP_TO_BEST: (
        "Score der semantisch besten URL der Query minus Score dieser URL. 0 bei der besten URL selbst; höchstens der "
        "Abstand 'fast gleich' aus Schritt 7c heißt praktisch gleichauf. Leer, wenn die URL nicht im Frog-Export steht."
    ),
    L.C_COMP_CHUNK: (
        "Der Textblock dieser URL, der zur Query am besten passt. Fehlt bei der Bewertungsgrundlage Gesamt-URL; leer, "
        "wenn die URL nicht im Frog-Export steht."
    ),
    L.C_COMP_POS: "Eigene Position dieser URL für die Query. Leer, wenn sie dafür nicht rankt.",
    # Potentielle Content-Lücken
    L.C_BEST_SCORE: "Score der besten URL nach der gewählten Bewertungsgrundlage. Er liegt unter der Schwelle aus Schritt 7b.",
    L.C_TOPIC: (
        "Nummer des SERP-Clusters aus Schritt 6: Queries mit gleicher Nummer haben stark überlappende Google-Ergebnisse "
        "und ergeben zusammen eine neue Seite. Leer ohne SERPs oder wenn die Query in keinem Cluster ist."
    ),
}

# "passt klar": eine Seite erreicht die Schwelle und liegt mehr als den Abstand 'fast gleich' vor jeder anderen passenden
VERDICT_HELP = {
    L.V_MATCH: (
        "Genau eine Seite passt klar: Sie erreicht die Schwelle und liegt deutlich vor allen anderen. Die Query rankt "
        "schwach, gar nicht, oder es sind keine Rankings geladen. Diese Seite ausbauen und intern stärken statt neu bauen."
    ),
    L.V_GAP: (
        "Keine Seite erreicht die Schwelle, und keine eigene Seite rankt bis zur Position aus "
        "luecke_nur_ohne_ranking_bis_position (Voreinstellung 20), oder es sind keine Rankings geladen."
    ),
    L.V_OK: (
        "Die Query rankt gut, und die rankende Seite erreicht die Schwelle. Steht die rankende URL nicht im Frog-Export, "
        "kann das Tool sie nicht prüfen und wertet ebenfalls 'In Ordnung'. Passt eine andere eigene Seite besser oder "
        "fast gleich gut, steht die Query zusätzlich im Blatt Kannibalisierungsgefahr: Heute rankt die richtige Seite, Google kann "
        "aber wechseln."
    ),
    L.V_CANNIBAL: (
        "Mehrere eigene Seiten erreichen die Schwelle und passen fast gleich gut, und die Query rankt schwach, gar nicht, "
        "oder es sind keine Rankings geladen. Details im Blatt Kannibalisierungsgefahr."
    ),
    L.V_WATCH: (
        "Die Query rankt gut, aber die rankende Seite erreicht die Schwelle nicht. Passt eine andere eigene Seite, steht "
        "die Query zusätzlich im Blatt Kannibalisierungsgefahr: prüfen, welche Seite die richtige ist. Passt gar keine Seite, die "
        "Schwelle prüfen. Ebenso, wenn keine Seite passt, aber eine eigene Seite bis zur Position aus "
        "luecke_nur_ohne_ranking_bis_position rankt (Voreinstellung 20): Google hält sie für relevant, also ausbauen "
        "statt neu bauen."
    ),
}

_FILLS = {
    L.V_MATCH: "FFEB9C",
    L.V_OK: "C6EFCE",
    L.V_CANNIBAL: "F8CBAD",
    L.V_GAP: "FFC7CE",
    L.V_WATCH: "D9D9D9",
}

DISCLAIMER = (
    "Das Notebook sortiert vor und begründet. Cosinus-Werte sind Hinweise, keine Urteile: "
    "die Entscheidung trifft ein Mensch nach Prüfung der Seite."
)

CAVEAT_SCORES = (
    "Scores eines Modells liegen in einem engen Band (bei e5 etwa 0,7 bis 0,9). "
    "Sie sind nur innerhalb eines Laufs vergleichbar, nicht zwischen Modellen."
)
# Herkunft der Schwelle -> fester Hinweis in der Lesehilfe
CAVEAT_THRESHOLD = {
    "median": (
        "Schwelle nicht kalibriert: Sie ist der mittlere beste Score, etwa die Hälfte der Queries liegt per "
        "Konstruktion darunter. Urteile und die Zahl neuer Seiten sind nur mit selbst geprüfter Schwelle belastbar."
    ),
    "rankings": (
        "Schwelle aus Rankings kalibriert: Sie ist so gewählt, dass 75 % der gut rankenden Paare sie erreichen. "
        "Bis zu ein Viertel der gut rankenden Queries erscheint daher als 'Rankt trotz schwachem Match'."
    ),
    "manuell": (
        "Schwelle von Hand gesetzt: Sie wurde im Notebook eingetragen und nicht aus den Daten abgeleitet. "
        "Urteile und die Zahl neuer Seiten hängen direkt von diesem Wert ab."
    ),
}


def _gap_sheet_help(new_pages) -> str:
    if new_pages is None:
        return f"{_SHEET_HELP[SHEET_GAPS]} Ohne SERPs (Schritt 6) gibt es kein Thema, jede Query zählt einzeln."
    return f"{_SHEET_HELP[SHEET_GAPS]} Neue Seiten: {new_pages} (eine je Thema, Queries ohne Thema einzeln)."


def _readme(sheets, settings, threshold_source, new_pages) -> pd.DataFrame:
    rows = [("Hinweis", "Einordnung", DISCLAIMER), ("Hinweis", "Scores", CAVEAT_SCORES)]
    if threshold_source in CAVEAT_THRESHOLD:
        rows.append(("Hinweis", "Schwelle", CAVEAT_THRESHOLD[threshold_source]))
    for name in list(sheets) + [SHEET_README]:
        text = _gap_sheet_help(new_pages) if name == SHEET_GAPS else _SHEET_HELP[name]
        rows.append(("Blatt", name, text))
    rows += [("Urteil", verdict, text) for verdict, text in VERDICT_HELP.items()]
    columns = dict.fromkeys(column for df in sheets.values() for column in df.columns)
    rows += [("Spalte", column, _COLUMN_HELP[column]) for column in columns]
    rows += [("Einstellung", key, str(value)) for key, value in settings.items()]
    return pd.DataFrame(rows, columns=[L.R_AREA, L.R_ENTRY, L.R_TEXT])


def build_sheets(overview, cannibal, gaps, settings, threshold_source=None, new_pages=None, elsewhere=None) -> dict:
    """threshold_source: "rankings", "median" oder "manuell" (Herkunft der Schwelle für die Lesehilfe).
    new_pages: Zahl neuer Seiten aus den Lücken, nur mit SERPs (sonst None)."""
    sheets = {
        SHEET_OVERVIEW: overview[OVERVIEW_COLUMNS],
        SHEET_CANNIBAL: cannibal[[c for c in CANNIBAL_COLUMNS if c in cannibal.columns]],  # ohne Chunk bei Gesamt-URL
        SHEET_GAPS: gaps[GAP_COLUMNS],
        SHEET_ELSEWHERE: (pd.DataFrame(columns=ELSEWHERE_COLUMNS) if elsewhere is None else elsewhere)[ELSEWHERE_COLUMNS],
    }
    return {**sheets, SHEET_README: _readme(sheets, settings, threshold_source, new_pages)}


def _clean(df: pd.DataFrame) -> pd.DataFrame:
    """Steuerzeichen entfernen, die openpyxl nicht schreiben kann. Das Original bleibt unverändert."""
    out = df.copy()
    for column in out.columns:
        if out[column].dtype == object or pd.api.types.is_string_dtype(out[column]):
            out[column] = out[column].map(lambda v: ILLEGAL_CHARACTERS_RE.sub(" ", v) if isinstance(v, str) else v)
    return out


_FORMULA_START = ("=", "+", "-", "@")


# Vorsprung vor zweitbester URL, gemessen am Abstand "fast gleich": höchstens der Abstand rot, bis zum Doppelten gelb
_LEAD_FILLS = ("F8CBAD", "FFEB9C", "C6EFCE")


def _lead_fill(value, margin):
    if value is None or pd.isna(value):
        return None
    if value <= margin + 1e-9:
        return _LEAD_FILLS[0]
    return _LEAD_FILLS[1] if value <= 2 * margin + 1e-9 else _LEAD_FILLS[2]


def write_excel(path, sheets: dict, margin=0.01) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, df in sheets.items():
            df = _clean(df)
            df.to_excel(writer, sheet_name=name, index=False)
            sheet = writer.sheets[name]
            for cell in sheet[1]:
                cell.font = Font(bold=True)
            for row in sheet.iter_rows():
                for cell in row:
                    # Text wie "=cmd|x" bleibt Text und wird nicht als Formel ausgeführt
                    if isinstance(cell.value, str) and cell.value.startswith(_FORMULA_START):
                        cell.data_type = "s"
            for idx, column in enumerate(df.columns, start=1):
                longest = max([len(str(column))] + [len(str(v)) for v in df[column].head(200)])
                sheet.column_dimensions[get_column_letter(idx)].width = min(max(12, longest + 2), 70)
            if L.C_VERDICT in df.columns:
                col = list(df.columns).index(L.C_VERDICT) + 1
                for row, verdict in enumerate(df[L.C_VERDICT], start=2):
                    colour = _FILLS.get(verdict)
                    if colour:
                        sheet.cell(row=row, column=col).fill = PatternFill("solid", fgColor=colour)
            if L.C_LEAD_GAP in df.columns:
                col = list(df.columns).index(L.C_LEAD_GAP) + 1
                for row, value in enumerate(df[L.C_LEAD_GAP], start=2):
                    colour = _lead_fill(value, margin)
                    if colour:
                        sheet.cell(row=row, column=col).fill = PatternFill("solid", fgColor=colour)


_PLAIN_NUMBER = re.compile(r"^-?\d+\.\d+$")
# Positionen liegen als Text vor ("4.3"), damit ganze Zahlen ohne ".0" erscheinen
_POSITION_COLUMNS = (L.C_POSITION, L.C_COMP_POS)


def _comma_in_text(name, df: pd.DataFrame) -> pd.DataFrame:
    """Zahlen, die als Text vorliegen, mit Dezimalkomma (sonst liest deutsches Excel "4.3" als Datum)."""
    out = df.copy()

    def swap(value):
        return value.replace(".", ",") if isinstance(value, str) else value

    for column in _POSITION_COLUMNS:
        if column in out.columns:
            out[column] = out[column].map(swap)
    if name == SHEET_README:
        numeric = (out[L.R_AREA] == "Einstellung") & out[L.R_TEXT].map(lambda v: bool(_PLAIN_NUMBER.match(str(v))))
        out.loc[numeric, L.R_TEXT] = out.loc[numeric, L.R_TEXT].map(swap)
    return out


def write_csv_zip(path, sheets: dict, sep: str = ";") -> None:
    # Semikolon ist das deutsche Excel-Format: dort gehört das Komma in die Zahl
    decimal = "," if sep == ";" else "."
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, df in sheets.items():
            if decimal == ",":
                df = _comma_in_text(name, df)
            archive.writestr(_CSV_NAMES[name], df.to_csv(index=False, sep=sep, decimal=decimal).encode("utf-8-sig"))
